from typing import Any
from src.core.domain.models import Document
from src.ports.outbound.document_reader_port import IDocumentReader

class TextPdfReader(IDocumentReader):
    """
    Extracts structured text/markdown from dense academic documents (papers, book chapters, lecture notes, presentations).
    Uses MarkItDown if available; falls back to PyMuPDF (fitz) for PDF files and python-pptx for PPTX files.
    Trade-off: Bypasses Gemini File API upload latency and remote file limits entirely,
    injecting extracted text directly into the LLM context.
    """
    def __init__(self):
        self._has_markitdown = False
        try:
            from markitdown import MarkItDown
            self._md = MarkItDown()
            self._has_markitdown = True
        except ImportError:
            self._has_markitdown = False

    def read(self, document: Document) -> str:
        doc_path = str(document.path)
        ext = document.path.suffix.lower()
        print(f"    [Local Parser] Extracting text/markdown from '{document.path.name}'...")

        if self._has_markitdown:
            try:
                result = self._md.convert(doc_path)
                text = result.text_content.strip()
                if text:
                    print(f"    [MarkItDown] Extracted {len(text)} characters.")
                    return text
            except Exception as exc:
                print(f"    [MarkItDown Warning] Conversion failed ({exc}), attempting fallback...")

        # Fallback to PyMuPDF (fitz) only for PDF files
        if ext == ".pdf":
            try:
                import fitz
                doc = fitz.open(doc_path)
                pages_text = []
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    txt = page.get_text("text").strip()
                    if txt:
                        pages_text.append(f"--- PAGE {page_num + 1} ---\n{txt}")
                doc.close()
                full_text = "\n\n".join(pages_text).strip()
                print(f"    [PyMuPDF] Extracted {len(full_text)} characters across {len(pages_text)} pages.")
                return full_text
            except Exception as exc:
                raise RuntimeError(f"Impossibile estrarre testo dal PDF '{doc_path}': {exc}")

        # Fallback to python-pptx for presentation files if MarkItDown fails
        if ext == ".pptx":
            try:
                from pptx import Presentation
                prs = Presentation(doc_path)
                slides_text = []
                for idx, slide in enumerate(prs.slides, 1):
                    slide_parts = [f"--- SLIDE {idx} ---"]
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                line = paragraph.text.strip()
                                if line:
                                    slide_parts.append(line)
                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                        notes = slide.notes_slide.notes_text_frame.text.strip()
                        if notes:
                            slide_parts.append(f"[Note: {notes}]")
                    if len(slide_parts) > 1:
                        slides_text.append("\n".join(slide_parts))
                full_text = "\n\n".join(slides_text).strip()
                if full_text:
                    print(f"    [python-pptx] Extracted {len(full_text)} characters across {len(slides_text)} slides.")
                    return full_text
            except Exception as exc:
                print(f"    [python-pptx Warning] Fallback failed ({exc}).")

        raise RuntimeError(f"Impossibile convertire il documento '{doc_path}' (formato {ext}).")

    def cleanup(self, payload: Any) -> None:
        # No remote resources to clean up for local text extraction
        pass
