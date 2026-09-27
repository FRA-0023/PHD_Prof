from pathlib import Path

def extract_pptx_to_markdown(path: Path) -> str:
    """
    Extracts structured Markdown from a PowerPoint (.pptx) file.
    Follows the markitdown-conversion skill pipeline:
    1. Attempts extraction via Microsoft's MarkItDown.
    2. Falls back to python-pptx for robust extraction of titles, shapes, tables, and speaker notes.
    """
    path_str = str(path)

    # 1. Primary: MarkItDown (official markdown converter)
    try:
        from markitdown import MarkItDown
        md = MarkItDown()
        result = md.convert(path_str)
        text = result.text_content.strip()
        if text:
            return text
    except Exception:
        pass

    # 2. Resilient Fallback: python-pptx (direct structural parsing)
    try:
        from pptx import Presentation
        prs = Presentation(path_str)
        slides_text = []
        for idx, slide in enumerate(prs.slides, 1):
            parts = [f"## Slide {idx}"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        line = paragraph.text.strip()
                        if line:
                            parts.append(line)
                elif shape.has_table:
                    table = shape.table
                    for row_idx, row in enumerate(table.rows):
                        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        parts.append("| " + " | ".join(cells) + " |")
                        if row_idx == 0:
                            parts.append("| " + " | ".join(["---"] * len(cells)) + " |")
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"> **Note del relatore**: {notes}")
            if len(parts) > 1:
                slides_text.append("\n\n".join(parts))

        full_text = "\n\n---\n\n".join(slides_text).strip()
        if full_text:
            return full_text
    except Exception as exc:
        raise RuntimeError(f"Impossibile estrarre testo dal file PPTX '{path_str}': {exc}")

    raise RuntimeError(f"Nessun contenuto estraibile trovato nel file PPTX '{path_str}'.")
